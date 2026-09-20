"""Export nutrition_repository_imputed.csv to a professionally formatted Excel (.xlsx) file."""

from __future__ import annotations

from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
CSV_PATH = ROOT / "data_processed" / "nutrition_repository_imputed.csv"
EXCEL_PATH = ROOT / "data_processed" / "nutrition_repository_imputed.xlsx"


def style_sheet(ws, is_summary: bool = False) -> None:
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    alt_fill = PatternFill(start_color="F2F7FA", end_color="F2F7FA", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Freeze panes (freeze header row + first 2 columns)
    ws.freeze_panes = "C2" if not is_summary else "B2"
    ws.auto_filter.ref = ws.dimensions

    # Style header
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28

    # Style data rows
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        ws.row_dimensions[row_idx].height = 20
        use_alt = row_idx % 2 == 0
        for col_idx, cell in enumerate(row, start=1):
            cell.border = thin_border
            if use_alt:
                cell.fill = alt_fill
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00" if isinstance(cell.value, float) else "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 40)


def main() -> None:
    print(f"[EXCEL] Loading {CSV_PATH.name}...")
    df = pd.read_csv(CSV_PATH)

    # Create Excel writer
    with pd.ExcelWriter(EXCEL_PATH, engine="openpyxl") as writer:
        # Sheet 1: Semua Data
        df.to_excel(writer, sheet_name="Semua Pangan (Lengkap)", index=False)

        # Sheet 2: Makronutrien
        macro_cols = [
            "food_id", "food_name_original", "source", "category_normalized",
            "water_g", "energy_kcal", "protein_g", "fat_g", "carbohydrate_g", "fiber_g", "ash_g"
        ]
        macro_df = df[[c for c in macro_cols if c in df.columns]]
        macro_df.to_excel(writer, sheet_name="Makronutrien & Energi", index=False)

        # Sheet 3: Mineral & Vitamin
        micro_cols = [
            "food_id", "food_name_original", "source", "category_normalized",
            "calcium_mg", "phosphorus_mg", "iron_mg", "sodium_mg", "potassium_mg", "copper_mg", "zinc_mg",
            "retinol_mcg", "beta_carotene_mcg", "carotene_total_mcg", "vitamin_a_mcg",
            "vitamin_b1_mg", "vitamin_b2_mg", "niacin_mg", "vitamin_c_mg"
        ]
        micro_df = df[[c for c in micro_cols if c in df.columns]]
        micro_df.to_excel(writer, sheet_name="Vitamin & Mineral", index=False)

        # Sheet 4: Ringkasan Kategori
        summary = (
            df.groupby(["source", "category_normalized"])
            .agg(
                Jumlah_Pangan=("food_id", "count"),
                Rata_Energi_kcal=("energy_kcal", "mean"),
                Rata_Protein_g=("protein_g", "mean"),
                Rata_Lemak_g=("fat_g", "mean"),
                Rata_Karbohidrat_g=("carbohydrate_g", "mean"),
                Rata_Serat_g=("fiber_g", "mean"),
            )
            .reset_index()
        )
        summary.to_excel(writer, sheet_name="Rekap per Kategori", index=False)

    # Load workbook to apply styles
    wb = openpyxl.load_workbook(EXCEL_PATH)
    style_sheet(wb["Semua Pangan (Lengkap)"])
    style_sheet(wb["Makronutrien & Energi"])
    style_sheet(wb["Vitamin & Mineral"])
    style_sheet(wb["Rekap per Kategori"], is_summary=True)
    wb.save(EXCEL_PATH)

    print(f"[EXCEL] Successfully generated {EXCEL_PATH} ({EXCEL_PATH.stat().st_size / 1024:.1f} KB)")


if __name__ == "__main__":
    main()
